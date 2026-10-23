# -*- coding: utf-8 -*-
"""Verify template upload -> placeholder detect -> fill -> generate."""
import os
import tempfile
from pathlib import Path

os.chdir(os.path.dirname(os.path.abspath(__file__)))
os.environ["TEACHFLOW_DB"] = os.path.join(tempfile.mkdtemp(), "t.db")
os.environ["TEACHFLOW_UPLOAD_ROOT"] = tempfile.mkdtemp()

from docx import Document  # noqa: E402
from app import app  # noqa: E402

tmp = Path(tempfile.mkdtemp())
tpl = tmp / "tpl.docx"
doc = Document()
doc.add_paragraph("Dear {{name}},")
doc.add_paragraph("Date: {{date}}")
doc.add_paragraph("Subject: {{subject}}")
doc.save(tpl)

c = app.test_client()
c.post("/register", data={"email": "t@x.com", "name": "T", "password": "secret123"})
c.post("/login", data={"email": "t@x.com", "password": "secret123"})

with open(tpl, "rb") as f:
    r = c.post("/studio/templates/upload", data={"name": "Letter", "file": (f, "tpl.docx")},
               content_type="multipart/form-data", follow_redirects=True)
print("upload:", r.status_code)
assert r.status_code == 200

r = c.get("/studio")
assert r.status_code == 200
body = r.data.decode("utf-8", "ignore")
print("placeholders shown on page:", "name" in body and "date" in body and "subject" in body)
assert "name" in body and "date" in body

r = c.post("/studio/templates/1/generate",
           data={"ph_name": "Ali", "ph_date": "01/01/2026", "ph_subject": "Meeting"},
           follow_redirects=True)
print("generate:", r.status_code, "| content-type:", r.headers.get("Content-Type", "")[:60])
assert r.status_code == 200

print("\nTEMPLATES OK")
