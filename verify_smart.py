# -*- coding: utf-8 -*-
"""Verify ChatGPT-style smart generation across document types."""
import os
import tempfile
from pathlib import Path

os.chdir(os.path.dirname(os.path.abspath(__file__)))

from smartgen import smart_generate  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

out = Path(tempfile.mkdtemp())

cases = [
    ("buat rph matematik tahun 5 topik pecahan", "docx", "Rancangan"),
    ("make an Excel marksheet for 30 students", "xlsx", "Markah"),
    ("buat excel kehadiran 25 murid", "xlsx", "Kehadiran"),
    ("jadual waktu kelas 6", "xlsx", "Jadual"),
    ("notis mesyuarat PIBG", "docx", "Notis"),
    ("presentation photosynthesis science", "pptx", ""),
    ("soalan uasa sains tahun 6", "docx", "Kuiz"),
    ("laporan aktiviti kokurikulum", "docx", "aktiviti"),
]
for text, kind, label in cases:
    r = smart_generate(text, out)
    f = out / r["fname"]
    assert f.exists() and f.stat().st_size > 0, text
    assert r["kind"] == kind, (text, r["kind"], r["label"])
    if label:
        assert label in r["label"], (text, r["label"])
    print(f"OK  {text!r:60} -> {r['kind']:5} {r['label']}")

# Tamil request
r = smart_generate("rph கணிதம் பின்னங்கள்", out)
assert r["kind"] == "docx"
print(f"OK  Tamil RPH -> {r['kind']} {r['label']}")

# Verify marksheet row count
r = smart_generate("buat excel markah matematik 30 murid", out)
wb = load_workbook(out / r["fname"])
ws = wb.active
assert ws.max_row == 31 and ws.max_column == 7, (ws.max_row, ws.max_column)
print(f"OK  marksheet 30 students -> {ws.max_row - 1} rows x {ws.max_column} cols")

print("\nSMART OK")
