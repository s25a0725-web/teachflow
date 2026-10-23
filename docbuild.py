# -*- coding: utf-8 -*-
"""Document Studio helpers: generate content, build DOCX/PPTX/PDF, convert, templates."""
import csv
import json
import os
import shutil
import subprocess
from pathlib import Path

BASE = Path(__file__).resolve().parent


# --------------------------------------------------------------------------
# AI generation (optional; falls back to offline templates)
# --------------------------------------------------------------------------
def _ai_generate_json(topic, doc_type, lang, instructions):
    key = os.environ.get("OPENAI_API_KEY") or os.environ.get("TEACHFLOW_AI_KEY")
    if not key:
        return None
    base = os.environ.get("OPENAI_BASE_URL", "https://api.openai.com/v1")
    model = os.environ.get("OPENAI_MODEL", "gpt-4o-mini")
    lang_name = {"en": "English", "ms": "Bahasa Malaysia", "ta": "Tamil"}.get(lang, "English")
    system = (
        "You are the AI Smart Teacher Assistant for SJK(T) Sepang. Generate a structured school document. "
        "Output ONLY valid JSON, no commentary, with this exact shape: "
        '{"title": "...", "sections": [{"heading": "...", "bullets": ["...", "..."]}]}. '
        "Write all content in " + lang_name + ". Include 3-5 sections, 2-4 bullets each."
    )
    user = f"Topic: {topic}\nDocument type: {doc_type}\nInstructions: {instructions or 'None'}"
    payload = {"model": model, "messages": [
        {"role": "system", "content": system}, {"role": "user", "content": user}], "max_tokens": 900}
    try:
        import urllib.request
        req = urllib.request.Request(
            base.rstrip("/") + "/chat/completions",
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json", "Authorization": "Bearer " + key})
        with urllib.request.urlopen(req, timeout=40) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        text = data["choices"][0]["message"]["content"].strip()
        text = text.strip("`")
        if text.startswith("json"):
            text = text[4:].strip()
        return json.loads(text)
    except Exception:
        return None


def generate_content(topic, doc_type, lang, instructions):
    ai = _ai_generate_json(topic, doc_type, lang, instructions)
    if ai and isinstance(ai, dict) and ai.get("sections"):
        return ai
    return _offline_content(topic, lang)


def _offline_content(topic, lang):
    T = {
        "en": {"i": "Introduction", "k": "Key Points", "d": "Details", "a": "Action Items",
               "bi": ["What is %s?", "Why it matters."],
               "bk": ["Main idea of %s.", "Important supporting facts."],
               "bd": ["Further explanation of %s.", "Examples and context."],
               "ba": ["Review the material on %s.", "Prepare a short summary."]},
        "ms": {"i": "Pengenalan", "k": "Perkara Utama", "d": "Butiran", "a": "Tindakan",
               "bi": ["Apakah itu %s?", "Mengapa ia penting."],
               "bk": ["Idea utama %s.", "Fakta sokongan yang penting."],
               "bd": ["Penerangan lanjut tentang %s.", "Contoh dan konteks."],
               "ba": ["Semak semula bahan tentang %s.", "Sediakan ringkasan ringkas."]},
        "ta": {"i": "அறிமுகம்", "k": "முக்கிய கருத்துகள்", "d": "விவரங்கள்", "a": "செயல்கள்",
               "bi": ["%s என்றால் என்ன?", "அது ஏன் முக்கியம்."],
               "bk": ["%s இன் முக்கிய யோசனை.", "முக்கிய ஆதரவு உண்மைகள்."],
               "bd": ["%s பற்றிய விளக்கம்.", "எடுத்துக்காட்டுகள்."],
               "ba": ["%s பற்றி மீள்பார்வை செய்யவும்.", "சிறு சுருக்கம் தயாரிக்கவும்."]},
    }.get(lang)
    if not T:
        T = _offline_content.__wrapped__ if False else None
    # safe fallback
    T = {
        "en": {"i": "Introduction", "k": "Key Points", "d": "Details", "a": "Action Items",
               "bi": ["What is %s?", "Why it matters."],
               "bk": ["Main idea of %s.", "Important supporting facts."],
               "bd": ["Further explanation of %s.", "Examples and context."],
               "ba": ["Review the material on %s.", "Prepare a short summary."]},
        "ms": {"i": "Pengenalan", "k": "Perkara Utama", "d": "Butiran", "a": "Tindakan",
               "bi": ["Apakah itu %s?", "Mengapa ia penting."],
               "bk": ["Idea utama %s.", "Fakta sokongan yang penting."],
               "bd": ["Penerangan lanjut tentang %s.", "Contoh dan konteks."],
               "ba": ["Semak semula bahan tentang %s.", "Sediakan ringkasan ringkas."]},
        "ta": {"i": "அறிமுகம்", "k": "முக்கிய கருத்துகள்", "d": "விவரங்கள்", "a": "செயல்கள்",
               "bi": ["%s என்றால் என்ன?", "அது ஏன் முக்கியம்."],
               "bk": ["%s இன் முக்கிய யோசனை.", "முக்கிய ஆதரவு உண்மைகள்."],
               "bd": ["%s பற்றிய விளக்கம்.", "எடுத்துக்காட்டுகள்."],
               "ba": ["%s பற்றி மீள்பார்வை செய்யவும்.", "சிறு சுருக்கம் தயாரிக்கவும்."]},
    }.get(lang, {
        "i": "Introduction", "k": "Key Points", "d": "Details", "a": "Action Items",
        "bi": ["What is %s?", "Why it matters."],
        "bk": ["Main idea of %s.", "Important supporting facts."],
        "bd": ["Further explanation of %s.", "Examples and context."],
        "ba": ["Review the material on %s.", "Prepare a short summary."]})
    return {
        "title": topic,
        "sections": [
            {"heading": T["i"], "bullets": [T["bi"][0] % topic, T["bi"][1]]},
            {"heading": T["k"], "bullets": [T["bk"][0] % topic, T["bk"][1]]},
            {"heading": T["d"], "bullets": [T["bd"][0] % topic, T["bd"][1]]},
            {"heading": T["a"], "bullets": [T["ba"][0] % topic, T["ba"][1]]},
        ],
    }


# --------------------------------------------------------------------------
# Fonts (Tamil support)
# --------------------------------------------------------------------------
def tamil_font():
    for p in ("C:/Windows/Fonts/Nirmala.ttf", "C:/Windows/Fonts/NirmalaUI.ttf",
              "C:/Windows/Fonts/Latha.ttf", "C:/Windows/Fonts/Nirmala.ttc"):
        if os.path.exists(p):
            return p
    return None


# --------------------------------------------------------------------------
# Builders
# --------------------------------------------------------------------------
def build_docx(content, out_path, meta=None):
    from docx import Document
    doc = Document()
    doc.add_heading(content.get("title") or "Document", 0)
    if meta:
        table = doc.add_table(rows=0, cols=2)
        for k, v in meta:
            row = table.add_row().cells
            row[0].text = str(k)
            row[1].text = str(v)
    for sec in content.get("sections", []):
        doc.add_heading(sec.get("heading", ""), level=1)
        for b in sec.get("bullets", []):
            doc.add_paragraph(b, style="List Bullet")
    doc.save(str(out_path))


def build_pptx(content, out_path):
    from pptx import Presentation
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = content.get("title") or "Document"
    for sec in content.get("sections", []):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = sec.get("heading", "")
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, b in enumerate(sec.get("bullets", [])):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = b
    prs.save(str(out_path))


def _esc(s):
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(content, out_path, meta=None):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_name = "Helvetica"
    ttf = tamil_font()
    if ttf:
        try:
            pdfmetrics.registerFont(TTFont("TamilFont", ttf))
            font_name = "TamilFont"
        except Exception:
            pass

    doc = SimpleDocTemplate(str(out_path), pagesize=A4,
                            leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=48)
    styles = {
        "title": ParagraphStyle("t", fontName=font_name, fontSize=20, leading=26, spaceAfter=14),
        "head": ParagraphStyle("h", fontName=font_name, fontSize=14, leading=18,
                               spaceBefore=12, spaceAfter=6, textColor="#1e293b"),
        "bullet": ParagraphStyle("b", fontName=font_name, fontSize=11, leading=16,
                                 leftIndent=16, bulletIndent=4, spaceAfter=4),
        "meta": ParagraphStyle("m", fontName=font_name, fontSize=10, leading=14, textColor="#475569"),
    }
    story = [Paragraph(_esc(content.get("title") or "Document"), styles["title"])]
    if meta:
        for k, v in meta:
            story.append(Paragraph("<b>%s:</b> %s" % (_esc(str(k)), _esc(str(v))), styles["meta"]))
        story.append(Paragraph("", styles["meta"]))
    for sec in content.get("sections", []):
        story.append(Paragraph(_esc(sec.get("heading", "")), styles["head"]))
        for b in sec.get("bullets", []):
            story.append(Paragraph("&bull; " + _esc(b), styles["bullet"]))
    doc.build(story)


# --------------------------------------------------------------------------
# Conversion
# --------------------------------------------------------------------------
def _soffice():
    for c in ("soffice", "C:/Program Files/LibreOffice/program/soffice.exe",
              "C:/Program Files (x86)/LibreOffice/program/soffice.exe"):
        if os.path.exists(c):
            return c
    return shutil.which("soffice")


def libreoffice_available():
    return _soffice() is not None


def _lo_convert(src, target, out_dir):
    soff = _soffice()
    if not soff:
        raise RuntimeError("LibreOffice is not installed on this computer.")
    subprocess.run([soff, "--headless", "--convert-to", target, "--outdir", str(out_dir), str(src)],
                   check=True, capture_output=True, timeout=120)


def convert_file(src_path, src_name, target):
    """Convert a file. Returns (output_name, mime). Raises RuntimeError if not possible."""
    src = Path(src_path)
    name = src_name.lower()
    stem = Path(src_name).stem or "document"
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    out_dir = src.parent

    if target == "pdf":
        if ext == "pdf":
            return src.name, "application/pdf"
        if ext in ("docx", "pptx", "xlsx", "xls", "odt", "ods"):
            _lo_convert(src, "pdf", out_dir)
            out = out_dir / (stem + ".pdf")
            if out.exists():
                return out.name, "application/pdf"
            raise RuntimeError("PDF conversion failed.")
        if ext in ("png", "jpg", "jpeg", "webp", "bmp"):
            from file_assistant import image_to_pdf
            out = out_dir / (stem + ".pdf")
            image_to_pdf(src, out)
            return out.name, "application/pdf"
        if ext in ("txt", "md", "csv"):
            from file_assistant import text_to_pdf_bytes
            out = out_dir / (stem + ".pdf")
            out.write_bytes(text_to_pdf_bytes(src.read_text(encoding="utf-8", errors="ignore")))
            return out.name, "application/pdf"

    if target == "docx":
        if ext == "docx":
            return src.name, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if ext == "pdf":
            try:
                from pdf2docx import Converter
                out = out_dir / (stem + ".docx")
                cv = Converter(str(src))
                cv.convert(str(out))
                cv.close()
                return out.name, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            except Exception:
                pass
            # fallback: text -> docx
            from pypdf import PdfReader
            text = "\n".join((p.extract_text() or "") for p in PdfReader(str(src)).pages)
            out = out_dir / (stem + ".docx")
            build_docx({"title": stem, "sections": [{"heading": "", "bullets": text.split("\n")}]}, out)
            return out.name, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    if target == "csv":
        if ext in ("xlsx", "xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(src, read_only=True)
            ws = wb.active
            out = out_dir / (stem + ".csv")
            with open(out, "w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                for row in ws.iter_rows(values_only=True):
                    w.writerow(row)
            return out.name, "text/csv"

    raise RuntimeError("This conversion is not supported for that file type.")


# --------------------------------------------------------------------------
# Templates ({{placeholder}} replacement)
# --------------------------------------------------------------------------
def _replace_text(text, data):
    for k, v in data.items():
        text = text.replace("{{" + k + "}}", str(v)).replace("{{ " + k + " }}", str(v))
    return text


def _fill_docx_runs(p, data):
    if not p.runs:
        return
    full = "".join(r.text for r in p.runs)
    new = _replace_text(full, data)
    if new != full:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""


def fill_docx_template(template_path, data, out_path):
    from docx import Document
    doc = Document(template_path)
    for p in doc.paragraphs:
        _fill_docx_runs(p, data)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _fill_docx_runs(p, data)
    doc.save(str(out_path))


def fill_pptx_template(template_path, data, out_path):
    from pptx import Presentation
    prs = Presentation(template_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        run.text = _replace_text(run.text, data)
    prs.save(str(out_path))


def fill_template(template_path, data, out_path):
    n = Path(template_path).name.lower()
    if n.endswith(".pptx"):
        fill_pptx_template(template_path, data, out_path)
    else:
        fill_docx_template(template_path, data, out_path)


def extract_placeholders(path):
    import re
    text = ""
    n = Path(path).name.lower()
    if n.endswith(".pptx"):
        from pptx import Presentation
        prs = Presentation(path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text += shape.text_frame.text + "\n"
    else:
        from docx import Document
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
    return sorted(set(re.findall(r"\{\{\s*([\w ]+?)\s*\}\}", text)))


def content_to_text(content):
    lines = [content.get("title", "")]
    for sec in content.get("sections", []):
        lines.append("## " + sec.get("heading", ""))
        for b in sec.get("bullets", []):
            lines.append("- " + b)
    return "\n".join(lines)


def parse_content_text(text):
    lines = (text or "").split("\n")
    title = lines[0].strip() if lines else "Document"
    sections = []
    cur = None
    for ln in lines[1:]:
        s = ln.rstrip()
        if s.startswith("## "):
            cur = {"heading": s[3:].strip(), "bullets": []}
            sections.append(cur)
        elif s.startswith("- ") and cur is not None:
            cur["bullets"].append(s[2:].strip())
        elif s.strip() and cur is not None:
            cur["bullets"].append(s.strip())
    return {"title": title, "sections": sections}


def build_xlsx(sheet_name, headers, rows, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Sheet1")[:31] or "Sheet1"
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    wb.save(str(out_path))


def ai_revise(text, instruction, lang):
    key = os.environ.get("OPENAI_API_KEY") or os.environ.get("TEACHFLOW_AI_KEY")
    if not key:
        return None
    base = os.environ.get("OPENAI_BASE_URL", "https://api.openai.com/v1")
    model = os.environ.get("OPENAI_MODEL", "gpt-4o-mini")
    lang_name = {"en": "English", "ms": "Bahasa Malaysia", "ta": "Tamil"}.get(lang, "English")
    system = ("You are a friendly assistant for teachers. Revise the given text exactly as instructed. "
              "Output ONLY the revised text, nothing else. Write in " + lang_name + ".")
    payload = {"model": model, "messages": [
        {"role": "system", "content": system},
        {"role": "user", "content": f"Instruction: {instruction}\n\nText:\n{text}"}], "max_tokens": 1200}
    try:
        import urllib.request
        req = urllib.request.Request(
            base.rstrip("/") + "/chat/completions",
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json", "Authorization": "Bearer " + key})
        with urllib.request.urlopen(req, timeout=40) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        return data["choices"][0]["message"]["content"].strip()
    except Exception:
        return None
